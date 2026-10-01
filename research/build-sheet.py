#!/usr/bin/env python3
"""Build the daily keyword research .xlsx from research/_data.json.

Usage: python3 research/build-sheet.py [YYYY-MM-DD]
Writes research/keywords-<date>.xlsx
"""
import json
import sys
import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
date = sys.argv[1] if len(sys.argv) > 1 else datetime.date.today().isoformat()

COLUMNS = [
    ("#", 5),
    ("query", 34),
    ("est_volume", 14),
    ("intent", 14),
    ("page_type", 13),
    ("category", 16),
    ("matched_service", 20),
    ("service_url", 30),
    ("affiliate?", 11),
    ("network", 14),
    ("commission", 24),
    ("cookie", 14),
    ("apply_url", 34),
    ("notes", 40),
    ("status", 13),
    ("your_affiliate_link", 34),
]

data = json.loads((HERE / "_data.json").read_text())
rows = data["rows"]

wb = Workbook()
ws = wb.active
ws.title = "keywords"

head_fill = PatternFill("solid", fgColor="4F46E5")
head_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="E2DED6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical="top", wrap_text=True)

# Header
for ci, (name, width) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=ci, value=name)
    c.fill = head_fill
    c.font = head_font
    c.alignment = Alignment(vertical="center")
    c.border = border
    ws.column_dimensions[c.column_letter].width = width
ws.row_dimensions[1].height = 22

# Rows
aff_fill_yes = PatternFill("solid", fgColor="E7F7EC")
aff_fill_unknown = PatternFill("solid", fgColor="FFF6E5")
for ri, r in enumerate(rows, start=2):
    r.setdefault("#", ri - 1)
    r.setdefault("status", "pending")
    r.setdefault("your_affiliate_link", "")
    for ci, (name, _w) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=ri, column=ci, value=r.get(name, ""))
        c.alignment = wrap
        c.border = border
    # tint the affiliate? cell
    aff = str(r.get("affiliate?", "")).lower()
    acell = ws.cell(row=ri, column=9)
    if aff.startswith("yes"):
        acell.fill = aff_fill_yes
    elif aff.startswith("unknown") or aff == "":
        acell.fill = aff_fill_unknown

# Freeze header, enable autofilter
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{len(rows)+1}"

# status dropdown
dv = DataValidation(type="list", formula1='"pending,approved,registered,rejected"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"O2:O{len(rows)+1}")

out = HERE / f"keywords-{date}.xlsx"
wb.save(out)
print(f"Wrote {out} with {len(rows)} rows")
